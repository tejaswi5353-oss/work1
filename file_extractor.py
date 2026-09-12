import csv
import io
import json
from pathlib import Path
from fastapi import HTTPException
import docx
import openpyxl
import PyPDF2

ALLOWED = {".pdf", ".csv", ".json", ".docx", ".xlsx", ".txt", ".log", ".md"}

def extract_text(content: bytes, filename: str) -> tuple[str, dict]:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext or filename}")

    page_count, row_count = None, None
    try:
        if ext == ".pdf":
            reader = PyPDF2.PdfReader(io.BytesIO(content))
            pages = [p.extract_text() or "" for p in reader.pages]
            text = "\n".join(pages)
            page_count = len(reader.pages)
        elif ext == ".csv":
            raw_text = content.decode("utf-8", errors="ignore")
            rows = [r for r in csv.reader(io.StringIO(raw_text)) if any(c.strip() for c in r)]
            if len(rows) > 1:
                headers, data_rows = rows[0], rows[1:]
                row_count = len(data_rows)
                text = "\n".join(", ".join(f"{h}: {val}" for h, val in zip(headers, r)) for r in data_rows)
            elif len(rows) == 1:
                row_count, text = 1, ", ".join(rows[0])
            else:
                row_count, text = 0, ""
        elif ext == ".json":
            parsed = json.loads(content.decode("utf-8", errors="ignore"))
            text = json.dumps(parsed, indent=2)
        elif ext == ".docx":
            doc = docx.Document(io.BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs if p.text)
        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            lines, total_rows = [], 0
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if cells:
                        total_rows += 1
                        lines.append(" | ".join(cells))
            wb.close()
            text, row_count = "\n".join(lines), total_rows
        else:  # .txt, .log, .md
            text = content.decode("utf-8", errors="ignore")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to extract {filename}: {str(e)}")

    truncated = len(text) > 8000
    final_text = text[:8000]
    meta = {
        "type": ext.lstrip(".").upper(),
        "size_kb": round(len(content) / 1024, 2),
        "char_count": len(final_text),
        "truncated": truncated,
    }
    if page_count is not None:
        meta["page_count"] = page_count
    if row_count is not None:
        meta["row_count"] = row_count

    return final_text, meta
